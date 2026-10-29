"""Builds the NidhiSetu SIH 2026 idea deck (PS SIH26265) on the official SIH 2026 template.

Keeps the template's logos, footer, slide titles and every "idea details pointer" word for word
(template rule 5), removes the instructions slide (template note), and fills each slide with
live prototype screenshots, the draw.io flowchart, measured results and sourced facts.

Run:  ../backend/.venv/bin/python build_deck.py   ->  NidhiSetu_SIH26265_Idea.pptx
"""
import json
import re

from PIL import Image as _I
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck_helpers import (BAD, BG, BLUE, INK, LINE, MUTED, NAVY, OK, WARN, WHITE, box, circle, delete_slide, notes,
                          pointer, remove_shape, set_runs_text, set_team_oval, set_title, shape_by_name, table, text)

TEMPLATE = "SIH2026-official-template.pptx"
OUT = "NidhiSetu_SIH26265_Idea.pptx"
TEAM_NAME = "[Team Name]"   # fill before export: must not contain the college name
TEAM_ID = "[Team ID]"
EVAL = json.load(open("../backend/data/eval/results.json"))
SIM = EVAL["closure_simulation"]
SAFFRON = "B8640F"
LADDER_RATIO = (lambda im: im.height / im.width)(_I.open("img/ladder.png"))
F, G, CAUGHT, FA = EVAL["forged"], EVAL["genuine"], EVAL["caught"], EVAL["false_alarms"]
PHOTOS = next(g["n"] for g in EVAL["groups"] if g["group"] == "genuine_photo")
SEC = f"{EVAL['median_ms_pdf'] / 1000:.2f} s"


def rich(s, size=10, color=INK, accent=BLUE, bullet=True, space_after=3):
    """'**bold**' segments become bold accent text: the one keyword the evaluator should see in each line."""
    runs = []
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", s)):
        if part:
            runs.append((part, {"bold": True, "color": accent} if i % 2 else {}))
    return {"runs": runs, "bullet": bullet, "space_after": space_after, "size": size, "color": color}


prs = Presentation(TEMPLATE)
delete_slide(prs, 6)  # "Important instructions" slide; the template says it can be deleted
s1, s2, s3, s4, s5, s6 = prs.slides
for s in (s2, s3, s4, s5, s6):
    set_team_oval(s)

# ================================================================ 1. TITLE PAGE
sub = shape_by_name(s1, "Subtitle")
for p in sub.text_frame.paragraphs:
    if p.runs:
        set_runs_text(p, "NidhiSetu")
        break
tb = shape_by_name(s1, "TextBox 9")
values = ["Problem Statement ID – SIH26265",
          "Problem Statement Title – Streamline of security deposit/ FDR portal.",
          "Theme – Smart Automation", "PS Category – Software", f"Team ID – {TEAM_ID}",
          f"Team Name (Registered on portal) – {TEAM_NAME}"]
for p, v in zip([p for p in tb.text_frame.paragraphs if p.text.strip()], values):
    set_runs_text(p, v)
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.05
    p.space_after = Pt(10)
    for r in p.runs:
        r.font.size = Pt(17)
notes(s1, "One line per member. Say: PS SIH26265 from AICTE's Regulation Bureau, Smart Automation, software.")

# ================================================================ 2. IDEA TITLE
set_title(s2, "NidhiSetu")
remove_shape(shape_by_name(s2, "TextBox 8"))
text(s2, 1.9, 0.98, 8.7, 0.32, [{"runs": [("Every security deposit confirmed with the bank, and tracked from submission to release",
                                             {"italic": True, "color": NAVY, "size": 12.5})], "align": PP_ALIGN.CENTER}])
pointer(s2, 0.4, 1.34, 8.9, "❖ Proposed Solution (Describe your Idea/Solution/Prototype)", size=14)

cols = [
    ("Detailed explanation of the proposed solution", 2.95, [
        "Reads each FDR or e-FDR: amount, bank, **IFSC**, dates, lien; **OCR** for phone photos",
        "Checks the **deposit rules**: in favour of AICTE, lien, amount, tenure, branch",
        "**Asks the bank**: e-FDR signature, term-deposit API, or a signed branch link",
        "**Tracks the lifecycle**: reminders at 90/60/30/7 days, escalation, renewal",
        "**Release workflow**: conditions checked, two officials, discharge letter to the bank",
    ]),
    ("How it addresses the problem", 2.95, [
        "A deposit today is a paper receipt checked once. In Chandigarh, **₹116.84 crore** of FDRs the bank had never issued went unnoticed **for over a year** (Tribune, Mar 2026)",
        "NidhiSetu shows **what the institution says next to what the bank confirms**, for every FDR, on day one",
        f"A **bank re-check every 90 days** finds an early closure in about **{SIM['ours_median_days']} days**, not at maturity",
    ]),
    ("Innovation and uniqueness of the solution", 2.75, [
        "**Bank-confirmed, not paper-trusted**: 3 tiers; never uses a QR or phone printed on the receipt",
        "**Continuous re-confirmation**, so a closed deposit cannot hide",
        "**Pledge registry**: one FDR cannot secure two institutions",
        "**Tamper evidence**: white patches, font changes, maturity-value arithmetic",
        "**Maker-checker release**, hash-chained audit trail",
    ]),
]
x = 0.4
for hdr, w, items in cols:
    box(s2, x, 1.72, w, 3.13, fill=BG, line=LINE)
    text(s2, x + 0.08, 1.77, w - 0.16, 0.48, [{"runs": [(hdr, {"bold": True, "color": NAVY, "size": 11.5})]}])
    text(s2, x + 0.08, 2.26, w - 0.16, 2.55, [rich(t, size=10) for t in items])
    x += w + 0.12
_ev = _I.open("img/shot_evidence.png")
_evw = min(3.51, 2.3 * _ev.width / _ev.height)
s2.shapes.add_picture("img/shot_evidence.png", Inches(9.42 + (3.51 - _evw) / 2), Inches(1.34), width=Inches(_evw))
text(s2, 9.42, 3.68, 3.51, 0.25, [{"runs": [("Working prototype: the receipt says ₹25 lakh, the bank says ₹2.5 lakh; numbered boxes mark the evidence",
                                              {"italic": True, "color": MUTED, "size": 8.5})]}])
box(s2, 9.42, 3.98, 3.51, 0.87, fill="EAF6F0", line="BFE3D1")
text(s2, 9.5, 4.02, 3.35, 0.8, [
    {"runs": [(f"{CAUGHT}/{F}", {"bold": True, "color": OK, "size": 17}), ("  forgeries caught   ", {"size": 9.5, "color": INK}),
              (f"{FA}/{G}", {"bold": True, "color": OK, "size": 17}), ("  false alarms", {"size": 9.5})]},
    {"runs": [(f"{SEC} per receipt on a laptop CPU · {EVAL['documents']} labelled test receipts", {"size": 8.5, "color": MUTED})]},
], anchor=MSO_ANCHOR.MIDDLE)
s2.shapes.add_picture("img/pipeline.png", Inches(0.4), Inches(5.0), width=Inches(12.53))
text(s2, 0.4, 6.4, 12.53, 0.3, [{"runs": [
    ("Demo video: ", {"bold": True, "color": NAVY}), ("[add unlisted link]", {"color": MUTED}),
    ("     Live prototype: ", {"bold": True, "color": NAVY}), ("[add link]", {"color": MUTED}),
    ("     Code: ", {"bold": True, "color": NAVY}), ("[add GitHub link]", {"color": MUTED})]}], size=10)
notes(s2, "30 seconds. Open with Chandigarh: ₹116.84 crore of FDRs that the bank never issued, and nobody noticed for a year. "
          "NidhiSetu asks the bank on day one and again every quarter. Point at the screenshot: receipt ₹25 lakh, bank ₹2.5 lakh, "
          "the white patch is box 1. Point at the green box: measured, not promised. Close: the officer decides; the bank's record decides what is true.")

# ================================================================ 3. TECHNICAL APPROACH
remove_shape(shape_by_name(s3, "TextBox 8"))
pointer(s3, 0.4, 1.22, 4.75, "Technologies to be used (e.g. programming languages, frameworks, hardware)", size=11.5, h=0.45)
stack = [
    "**FastAPI · Python 3.12**: typed API; Python owns the PDF, OCR and signature libraries (vs Node.js: weak PDF tooling)",
    "**PostgreSQL 16**: every instrument with paper and bank values side by side; demo runs the same code on SQLite (vs MongoDB: weak transactions for financial records)",
    "**Redis + RQ workers**: the daily monitor and bank re-checks run as jobs (vs Celery: heavier; Kafka: overkill)",
    "**MinIO (S3) on NIC / MeghRaj**: receipts stored by SHA-256, encrypted (vs local disk: does not scale)",
    "**pyHanko**: verifies the bank's PAdES signature on an e-FDR; catches edits after signing",
    "**RapidOCR (PP-OCR, ONNX)**: reads phone photos on a CPU (vs cloud OCR: financial data leaves AICTE)",
    "**Bank channel adapters**: e-FDR, term-deposit API, signed branch link; Account Aggregator next (vs letters to branches: weeks, no trail)",
    "**RBI IFSC list + rapidfuzz**: the branch must exist; OCR-tolerant name matching (vs exact match: photos fail)",
    "**React 18 + TypeScript on UX4G 3.0**: the Government of India design system, accessibility bar (vs a custom look)",
    "**Role-based sign-in + SHA-256 hash chain**: maker-checker roles; tamper-evident log (vs blockchain: same proof, no cost)",
]
text(s3, 0.4, 1.72, 4.75, 4.45, [rich(t, size=9.8, space_after=4.5) for t in stack])
box(s3, 0.4, 6.25, 4.75, 0.5, fill=OK, line=None, radius=0.2)
text(s3, 0.45, 6.27, 4.65, 0.46, [{"runs": [("100% open source  ·  ₹0 licence  ·  CPU only, no GPU", {"bold": True, "color": WHITE, "size": 11.5})],
                                  "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
pointer(s3, 5.35, 1.22, 7.58, "Methodology and process for implementation (Flow Charts/Images/ working prototype)", size=11.5)
_fl = _I.open("img/flow.png")
_fh = min(4.62, 7.58 * _fl.height / _fl.width)
_fw = _fh * _fl.width / _fl.height
s3.shapes.add_picture("img/flow.png", Inches(5.35 + (7.58 - _fw) / 2), Inches(1.55), width=Inches(_fw))  # built in draw.io: img-src/flow.drawio
steps = [("1  Built", "lifecycle portal, bank sandbox, monitor, 6 passing tests"),
         ("2  Measured", f"{EVAL['documents']} labelled receipts, 8 forgery patterns"),
         ("3  Pilot", "our college's deposit records, then one regional office")]
sw = (7.58 - 2 * 0.1) / 3
for i, (h, d) in enumerate(steps):
    xx = 5.35 + i * (sw + 0.1)
    box(s3, xx, 6.25, sw, 0.5, fill="EEF0FC", line="DDE1FA")
    text(s3, xx + 0.08, 6.27, sw - 0.16, 0.46, [{"runs": [(h + "  ", {"bold": True, "color": NAVY, "size": 9.5}), (d, {"size": 8.5})]}],
         anchor=MSO_ANCHOR.MIDDLE)
notes(s3, "Left: every choice has a reason and a rejected alternative; ask us about any of them. "
          "Right, follow the arrows: the institution uploads the FDR; five engines read it, check the rules, ask the bank, look for "
          "tampering and decide. The lifecycle monitor runs every day: reminders, escalation, and a bank re-check every 90 days. "
          "The officer sees paper and bank side by side, asks the institution to explain a discrepancy, or starts a release that two "
          "different officials must approve before the bank is asked to remove AICTE's lien. The renewal loop on the left closes the cycle. "
          "The diagram is an editable draw.io file in our repo.")

# ================================================================ 4. FEASIBILITY AND VIABILITY
remove_shape(shape_by_name(s4, "TextBox 8"))
pointer(s4, 0.4, 1.22, 8.0, "Analysis of the feasibility of the idea")
tiles = [("Technical", "Built and running today; banks already issue e-receipts and share term deposits through AA", BLUE),
         ("Economic", "₹0 licence; replaces verification letters and follow-up calls; NIC cloud", OK),
         ("Operational", "Fits the Regulation Bureau's steps; officers keep every decision", NAVY),
         ("Legal & data", "On-prem, role-based, encrypted; bank checks only where authorised", SAFFRON)]
tw = (8.0 - 3 * 0.12) / 4
for i, (h, d, col) in enumerate(tiles):
    xx = 0.4 + i * (tw + 0.12)
    box(s4, xx, 1.55, tw, 0.95, fill=BG, line=LINE)
    text(s4, xx + 0.08, 1.58, tw - 0.16, 0.3, [{"runs": [(h, {"bold": True, "color": col, "size": 11.5})]}])
    text(s4, xx + 0.08, 1.87, tw - 0.16, 0.62, [d], size=9.3)
box(s4, 0.4, 2.6, 8.0, 1.05, fill="EAF6F0", line="BFE3D1")
nums = [(f"{CAUGHT}/{F}", "forgeries caught"), (f"{FA}/{G}", f"false alarms ({PHOTOS} phone photos)"),
        (SEC, "per receipt, laptop CPU"), (f"{SIM['ours_median_days']} days", f"to catch an early closure (paper: {SIM['paper_median_days']:,})")]
nw = 8.0 / 4
for i, (n, lab) in enumerate(nums):
    text(s4, 0.45 + i * nw, 2.62, nw - 0.1, 0.72, [{"runs": [(n, {"bold": True, "color": OK, "size": 20})]},
                                                   {"runs": [(lab, {"size": 9.2, "color": INK})]}])
text(s4, 0.5, 3.36, 7.8, 0.27, [{"runs": [(f"Measured on {EVAL['documents']} synthetic receipts (8 patterns from reported cases) and a 10,000-deposit "
                                            "simulation; the pilot measures real-world rates.", {"italic": True, "size": 8.5, "color": MUTED})]}])
pointer(s4, 0.4, 3.74, 3.4, "Potential challenges and risks", size=11.5)
pointer(s4, 3.85, 3.74, 4.6, "Strategies for overcoming these challenges", size=11.5)
risks = [["Challenge or risk", "Strategy"],
         ["Bank has no API (a small co-operative bank)", "Tier 3 signed branch link; no answer within the SLA escalates; AA consent as banks join"],
         ["Bank service down or slow", "Stays \"awaiting bank\", never \"fake\"; retried automatically"],
         ["Genuine receipt flagged (phone photo, OCR)", f"OCR-tolerant matching: 0/{PHOTOS} false alarms on photos; officer sees the evidence first"],
         ["FD closed after it was submitted", "Bank re-check every 90 days; lien must show at the bank, not only on paper"],
         ["A clever PDF edit", "The bank's record decides, not the file; forensics is a second signal"],
         ["Rules or amounts change", "Rulebook in YAML, versioned per cycle; no code change"]]
table(s4, 0.4, 4.03, 8.0, [3.2, 4.8], risks, size=9.2, row_h=0.39)
s4.shapes.add_picture("img/ladder.png", Inches(8.62), Inches(1.22), width=Inches(4.3))
LAD_H = 4.3 * LADDER_RATIO
box(s4, 8.62, 1.22 + LAD_H + 0.12, 4.3, 6.8 - (1.22 + LAD_H + 0.12), fill=BG, line=LINE)
text(s4, 8.72, 1.22 + LAD_H + 0.17, 4.1, 6.7 - (1.22 + LAD_H + 0.17), [
    {"runs": [("Viability", {"bold": True, "color": NAVY, "size": 12})]},
    rich("**Pilot**: our college's deposit records (with permission), then one AICTE regional office, then national", size=9.8),
    rich("**Reuse**: UGC, NCTE, PCI and any department that holds FDRs as security, like the municipal bodies in the 2026 cases", size=9.8),
    rich("**Sustainable**: open source under AICTE; each bank is onboarded once, through an adapter", size=9.8),
    rich("**Scale estimate**: if AICTE holds ~20,000 deposits, quarterly re-checks are ~220 bank calls a day; one worker does that in under a minute", size=9.8),
])
notes(s4, "Say the false-alarm answer first, judges always ask it: zero on 60 genuine receipts, including phone photos. "
          "Then the 45 days: that is what continuous re-confirmation buys. Right side: the three tiers explain how we confirm "
          "even when a bank has no API.")

# ================================================================ 5. IMPACT AND BENEFITS
remove_shape(shape_by_name(s5, "TextBox 8"))
pointer(s5, 0.4, 1.22, 6.1, "Potential impact on the target audience")
aud = [("R", "AICTE Regulation Bureau", "One register for every deposit; reminders and escalations automatic; release with a full trail", NAVY),
       ("I", "Honest institutions", "Upload once, get reminded before maturity, get the deposit released without chasing letters", OK),
       ("B", "Banks", "Answer once through an API or a one-click link, instead of letters and phone calls", SAFFRON),
       ("S", "Students and the public", "The security behind each approved institution is real money at a bank, not a piece of paper", BLUE)]
for i, (ic, h, d, col) in enumerate(aud):
    y = 1.6 + i * 0.86
    circle(s5, 0.45, y + 0.06, 0.5, col, ic, size=14)
    text(s5, 1.1, y, 5.4, 0.8, [{"runs": [(h, {"bold": True, "color": NAVY, "size": 12})]}, d], size=10.5)
pointer(s5, 6.75, 1.22, 6.2, "Benefits of the solution (social, economic, environmental, etc.)")
ben = [("Social", ["Deposits that back approvals are real", "Faster, transparent release for genuine institutions",
                   "Fewer disputes: every step on record"], BLUE),
       ("Economic", ["A fake or closed FDR is caught in weeks, not years", "Officer time saved on letters and follow-up",
                     "₹0 licence, no per-check fee"], OK),
       ("Environmental", ["Paperless from upload to discharge letter", "No couriered receipts or verification letters",
                          "Bank e-FDRs preferred over paper"], "2E7D32"),
       ("Governance", ["Institution's claim and bank's record kept apart, as the PS asks", "Two officials on every release",
                       "Tamper-evident audit trail"], NAVY)]
bw, bh = 3.02, 1.62
for i, (h, d, col) in enumerate(ben):
    xx = 6.75 + (i % 2) * (bw + 0.14)
    y = 1.58 + (i // 2) * (bh + 0.12)
    box(s5, xx, y, bw, bh, fill=BG, line=LINE)
    text(s5, xx + 0.1, y + 0.08, bw - 0.2, 0.3, [{"runs": [(h, {"bold": True, "color": col, "size": 12.5})]}])
    text(s5, xx + 0.1, y + 0.42, bw - 0.2, bh - 0.46, [{"runs": [(t, {})], "bullet": True, "space_after": 2} for t in d], size=10.5)
text(s5, 0.4, 5.1, 12.5, 0.3, [{"runs": [("Why it matters: reported cases", {"bold": True, "color": NAVY, "size": 12})]}])
stats = [("₹116.84 Cr", "of Chandigarh MC's FDRs were not in the bank's system; nobody noticed for over a year (The Tribune, Mar 2026)"),
         ("₹150 Cr", "of Panchkula MC's fixed deposits did not exist; found only on a maturity transfer request (Moneylife, Mar 2026)"),
         ("₹64 Cr", "swindled from Telugu Akademi's fixed deposits using forged FDRs and letters (Moneylife, 2021)"),
         ("IFN 760 COV", "Bank guarantees travel bank-to-bank on SFMS, the only valid mode (PNB). NidhiSetu gives FDRs a bank-side check")]
sw = (12.53 - 3 * 0.14) / 4
for i, (n, d) in enumerate(stats):
    xx = 0.4 + i * (sw + 0.14)
    box(s5, xx, 5.42, sw, 1.38, fill=WHITE, line=LINE)
    text(s5, xx + 0.1, 5.44, sw - 0.2, 0.5, [{"runs": [(n, {"bold": True, "color": BAD if i < 3 else NAVY, "size": 22 if i < 3 else 19})]}])
    text(s5, xx + 0.1, 5.95, sw - 0.2, 0.85, [d], size=9.3, color=INK)
notes(s5, "Who benefits, then the cases. Each number has its source on the card; do not round them up. "
          "The last card is the design idea: bank guarantees already need the bank's own message; we give FDRs the same.")

# ================================================================ 6. RESEARCH AND REFERENCES
remove_shape(shape_by_name(s6, "TextBox 8"))
pointer(s6, 0.4, 1.22, 12.5, "Details / Links of the reference and research work")
text(s6, 0.4, 1.55, 6.6, 0.3, [{"runs": [("How NidhiSetu compares with what exists today", {"bold": True, "color": NAVY, "size": 11.5})]}])
Y, N = ("Yes", OK), ("No", BAD)
cmp_rows = [["", "Paper file + letters", "Spreadsheet tracker", "NeSL e-BG", "Account Aggregator", "NidhiSetu"],
            ["Bank confirms the deposit", ("Letters", WARN), N, ("BGs only", WARN), ("With consent", WARN), Y],
            ["Covers FDRs held as security", Y, Y, N, ("Data only", WARN), Y],
            ["Reminders before maturity", N, ("Manual", WARN), ("n/a", MUTED), N, Y],
            ["Catches an early closure", ("At maturity", BAD), N, ("n/a", MUTED), ("If re-fetched", WARN), ("Every 90 d", OK)],
            ["Two-official release, audit trail", ("Partly", WARN), N, ("n/a", MUTED), N, Y],
            ["Paper vs bank, side by side", N, N, ("n/a", MUTED), N, Y]]
table(s6, 0.4, 1.9, 6.6, [1.85, 0.95, 0.95, 0.8, 1.05, 1.0], cmp_rows, size=9, row_h=0.4)
text(s6, 0.4, 4.85, 6.6, 1.95, [
    {"runs": [("What exists, and what we build on", {"bold": True, "color": NAVY, "size": 11.5})]},
    rich("**NeSL e-BG**: banks already issue bank guarantees electronically with instant verification; we bring that idea to FDRs", size=9.8),
    rich("**Account Aggregator**: term deposits are a listed data type; our next bank channel, with the institution's consent", size=9.8),
    rich("**RBI IFSC lists** (open dataset): every branch on a receipt must exist", size=9.8),
    {"runs": [("Stakeholder input: [add one line from your college's accounts officer or AICTE coordinator]", {"color": MUTED})], "bullet": True, "size": 9.8},
])
refs = [
    ("PS SIH26265, AICTE Regulation Bureau (sih.gov.in)", "https://sih.gov.in/sih2026PS"),
    ("AICTE, Simple steps for making security deposit online", "https://aicte.gov.in/sites/default/files/Security_Deposit.pdf"),
    ("The Tribune, Chandigarh MC: Rs 116 crore FDR scandal unfolds, 11 Mar 2026", "https://www.tribuneindia.com/news/chandigarh/chandigarh-mc-scam-rs-116-crore-fdr-scandal-unfolds-officials-under-fire/"),
    ("Moneylife, ₹150 crore FD discrepancy in Panchkula MC account, Mar 2026", "https://moneylife.in/article/150-crore-fd-discrepancy-in-panchkula-mc-account-kotak-mahindra-bank-files-police-complaint/80046.html"),
    ("Business Today, CBI arrests ex-bank officials in Haryana FD fraud case, 30 Jun 2026", "https://www.businesstoday.in/latest/corporate/story/cbi-arrests-former-idfc-first-bank-au-small-finance-bank-officials-in-haryana-fd-fraud-case-540047-2026-06-30"),
    ("Moneylife, ED to probe Rs64 crore Telugu Akademi fixed deposit fraud", "https://www.moneylife.in/article/ed-to-probe-rs64-crore-telugu-akademi-fixed-deposit-fraud/65336.html"),
    ("PNB, Bank guarantees transmitted through SFMS IFN 760 / 760 COV", "https://pnb.bank.in/working-capital-financing.html"),
    ("CSB Bank, Electronic bank guarantee through NeSL", "https://csb.bank.in/pdf/electronic-bank-guarantee-18082026.pdf"),
    ("Sahamati, Account types activated by banks on AAs (term deposits)", "https://sahamati.org.in/account-types-activated-by-banks-on-aas/"),
    ("Sahamati, Glossary of AA terms (TD: term / fixed deposit)", "https://sahamati.org.in/glossary-of-aa-terms/"),
    ("Razorpay IFSC dataset, built from RBI NEFT/RTGS lists", "https://github.com/razorpay/ifsc"),
    ("ETSI EN 319 142 (PAdES): PDF digital signatures", "https://www.etsi.org/"),
    ("Du et al., PP-OCR: a practical ultra lightweight OCR system, arXiv:2009.09941", "https://arxiv.org/abs/2009.09941"),
    ("UX4G 3.0, Government of India design system", "https://www.ux4g.gov.in/"),
    ("GIGW 3.0: Indian government web guidelines (WCAG 2.1 AA)", "https://guidelines.india.gov.in/"),
]
text(s6, 7.25, 1.55, 5.68, 0.3, [{"runs": [("Research, standards and sources", {"bold": True, "color": NAVY, "size": 11.5})]}])
text(s6, 7.25, 1.86, 5.68, 4.95, [{"runs": [(f"{i}. ", {"bold": True, "color": NAVY}), (t, {"link": u, "color": BLUE})], "space_after": 3}
                                  for i, (t, u) in enumerate(refs, 1)], size=8.8)
box(s6, 7.25, 5.0, 5.68, 1.8, fill=BG, line=LINE)
text(s6, 7.35, 5.05, 5.48, 1.72, [
    {"runs": [("Data and test assets (all free)", {"bold": True, "color": NAVY, "size": 11})]},
    rich(f"**Labelled test set we built**: {EVAL['documents']} receipts, 8 forgery patterns from the cases above", size=9.2),
    rich("**Sandbox banks**: our own CA for signed e-FDRs, a term-deposit API, branch links", size=9.2),
    rich("**10,000-deposit simulation** of early closures: paper check vs 90-day re-check", size=9.2),
    rich("**Rules**: AICTE security deposit norms, configurable per cycle (PS allows synthetic data)", size=9.2),
])
notes(s6, "If asked why not just use NeSL: it covers bank guarantees, not FDRs held as security. If asked about Account Aggregator: "
          "it is the right long-term channel and we have it as Tier 2 next; the institution consents once.")

prs.save(OUT)
print("saved", OUT)
